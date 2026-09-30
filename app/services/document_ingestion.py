"""Text extraction and deterministic chunking for uploaded documents."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from enum import StrEnum
from hashlib import sha256
from pathlib import PurePath
from uuid import UUID, uuid4

from sqlalchemy.orm import Session

from app.domain.documents import DocumentVersionStatus, EmbeddingSetStatus
from app.models.documents import (
    ChunkModel,
    DocumentModel,
    DocumentVersionModel,
    EmbeddingSetModel,
)
from app.providers.embeddings import EmbeddingProvider
from app.repositories.postgres.embeddings import PostgresEmbeddingRepository
from app.services.hr_chunking import HRDocumentChunk
from app.services.hr_embedding import embed_hr_chunks, persist_hr_chunk_embeddings

DEFAULT_MAX_CHUNK_CHARS = 1200
_PARAGRAPH_BOUNDARY_PATTERN = re.compile(r"\n[ \t]*\n+")
_SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx"})


class DocumentIngestionErrorCode(StrEnum):
    UNSUPPORTED_FORMAT = "unsupported_format"
    INVALID_DOCUMENT = "invalid_document"
    EMPTY_DOCUMENT = "empty_document"
    INVALID_MAX_CHUNK_CHARS = "invalid_max_chunk_chars"
    PARSER_UNAVAILABLE = "parser_unavailable"


class DocumentIngestionError(ValueError):
    """Raised when an uploaded document cannot be extracted or chunked."""

    def __init__(self, code: DocumentIngestionErrorCode, message: str) -> None:
        self.code = code
        super().__init__(message)


@dataclass(frozen=True, slots=True)
class DocumentChunk:
    """Extracted text chunk with a stable content digest."""

    chunk_index: int
    content_text: str
    content_hash: str


@dataclass(frozen=True, slots=True)
class IngestedDocument:
    """Normalized source text and chunks for an uploaded document."""

    filename: str
    extracted_text: str
    chunks: tuple[DocumentChunk, ...]


@dataclass(frozen=True, slots=True)
class PersistedIngestedDocument:
    """Database identity and initial lifecycle state for an uploaded document."""

    document_key: str
    version_id: UUID
    version_status: DocumentVersionStatus
    chunk_count: int


def _extract_text(filename: str, content: bytes) -> str:
    extension = "." + filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if extension not in _SUPPORTED_EXTENSIONS:
        raise DocumentIngestionError(
            DocumentIngestionErrorCode.UNSUPPORTED_FORMAT,
            "Upload a TXT, Markdown, PDF, or DOCX document.",
        )

    try:
        if extension in {".txt", ".md"}:
            return content.decode("utf-8-sig").strip()
        if extension == ".pdf":
            try:
                import fitz
            except ModuleNotFoundError as error:
                raise DocumentIngestionError(
                    DocumentIngestionErrorCode.PARSER_UNAVAILABLE,
                    "PDF reading is unavailable until PyMuPDF is installed.",
                ) from error
            pdf_document = None
            try:
                pdf_document = fitz.open(stream=content, filetype="pdf")
                return "\n".join(page.get_text() for page in pdf_document).strip()
            finally:
                if pdf_document is not None:
                    pdf_document.close()

        try:
            from docx import Document
        except ModuleNotFoundError as error:
            raise DocumentIngestionError(
                DocumentIngestionErrorCode.PARSER_UNAVAILABLE,
                "DOCX reading is unavailable until the document parser dependencies are installed.",
            ) from error
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        return "\n\n".join(paragraph for paragraph in paragraphs if paragraph)
    except DocumentIngestionError:
        raise
    except Exception as error:
        raise DocumentIngestionError(
            DocumentIngestionErrorCode.INVALID_DOCUMENT,
            "The document could not be read. Check that it is a valid, unencrypted file.",
        ) from error


def _split_oversized_text(text: str, max_chunk_chars: int) -> list[str]:
    chunks: list[str] = []
    remaining = text.strip()
    while len(remaining) > max_chunk_chars:
        split_at = remaining.rfind(" ", 0, max_chunk_chars + 1)
        if split_at <= 0:
            split_at = max_chunk_chars
        chunks.append(remaining[:split_at].strip())
        remaining = remaining[split_at:].strip()
    if remaining:
        chunks.append(remaining)
    return chunks


def _chunk_text(text: str, max_chunk_chars: int) -> tuple[DocumentChunk, ...]:
    paragraphs = [
        paragraph.strip()
        for paragraph in _PARAGRAPH_BOUNDARY_PATTERN.split(text.strip())
        if paragraph.strip()
    ]
    content_chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > max_chunk_chars:
            if current:
                content_chunks.append(current)
                current = ""
            content_chunks.extend(_split_oversized_text(paragraph, max_chunk_chars))
            continue

        candidate = paragraph if not current else f"{current}\n\n{paragraph}"
        if len(candidate) <= max_chunk_chars:
            current = candidate
        else:
            content_chunks.append(current)
            current = paragraph

    if current:
        content_chunks.append(current)

    return tuple(
        DocumentChunk(
            chunk_index=index,
            content_text=chunk,
            content_hash=sha256(chunk.encode("utf-8")).hexdigest(),
        )
        for index, chunk in enumerate(content_chunks)
    )


def ingest_document(
    filename: str,
    content: bytes,
    *,
    max_chunk_chars: int = DEFAULT_MAX_CHUNK_CHARS,
) -> IngestedDocument:
    """Extract supported document text and split it into stable chunks."""
    if max_chunk_chars <= 0:
        raise DocumentIngestionError(
            DocumentIngestionErrorCode.INVALID_MAX_CHUNK_CHARS,
            "max_chunk_chars must be greater than zero.",
        )

    extracted_text = _extract_text(filename, content)
    if not extracted_text:
        raise DocumentIngestionError(
            DocumentIngestionErrorCode.EMPTY_DOCUMENT,
            "No readable text was found in the document.",
        )

    return IngestedDocument(
        filename=filename,
        extracted_text=extracted_text,
        chunks=_chunk_text(extracted_text, max_chunk_chars),
    )


def persist_ingested_document(
    session: Session,
    *,
    tenant_id: str,
    content: bytes,
    document: IngestedDocument,
    embedding_provider: EmbeddingProvider,
    document_type: str = "Other",
) -> PersistedIngestedDocument:
    """Store extracted chunks and vectors as a non-searchable candidate version."""
    document_key = f"UPLOAD-{uuid4().hex.upper()}"
    document_model = DocumentModel(
        tenant_id=tenant_id,
        document_key=document_key,
        title=PurePath(document.filename).name[:300] or "Uploaded document",
        document_type=document_type,
    )
    session.add(document_model)
    session.flush()

    version = DocumentVersionModel(
        document_id=document_model.id,
        version_label="1.0",
        status=DocumentVersionStatus.CANDIDATE.value,
        content_hash=sha256(content).hexdigest(),
    )
    session.add(version)
    session.flush()

    chunk_models = [
        ChunkModel(
            document_version_id=version.id,
            chunk_index=chunk.chunk_index,
            heading_path=PurePath(document.filename).name[:500],
            content_text=chunk.content_text,
            content_hash=chunk.content_hash,
        )
        for chunk in document.chunks
    ]
    embedding_set = EmbeddingSetModel(
        document_version_id=version.id,
        model_name=embedding_provider.model_name,
        model_version=embedding_provider.model_version,
        dimension=embedding_provider.dimension,
        status=EmbeddingSetStatus.CANDIDATE.value,
    )
    session.add_all([*chunk_models, embedding_set])
    session.flush()

    hr_chunks = tuple(
        HRDocumentChunk(
            chunk_index=chunk.chunk_index,
            heading_path=PurePath(document.filename).name[:500],
            content_text=chunk.content_text,
            content_hash=chunk.content_hash,
        )
        for chunk in document.chunks
    )
    persist_hr_chunk_embeddings(
        PostgresEmbeddingRepository(session),
        embedding_set_id=embedding_set.id,
        embeddings=embed_hr_chunks(hr_chunks, embedding_provider),
        chunk_ids={model.chunk_index: model.id for model in chunk_models},
    )

    return PersistedIngestedDocument(
        document_key=document_key,
        version_id=version.id,
        version_status=DocumentVersionStatus.CANDIDATE,
        chunk_count=len(chunk_models),
    )